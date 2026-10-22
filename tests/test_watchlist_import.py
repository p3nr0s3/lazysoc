"""Tests for uploading / exporting the critical tools list (core.watchlist_import)."""

import io
import json

import pytest
from openpyxl import Workbook, load_workbook

import core.watchlist_import as wi


def _xlsx(sheets):
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _asset(watchlist, asset_id):
    return next(a for a in watchlist["assets"] if a["id"] == asset_id)


def test_template_csv_is_enriched_from_catalog():
    csv_text = (
        "owner,product,vendor,version,criticality,category,note,aliases,cpe,kev\n"
        "Internal,FortiGate,Fortinet,7.4.8,,Firewall,edge-fw-01,,,\n"
        "Customer A,Fortigate 100F,Fortinet,,,,,,,\n"
    )
    watchlist, report = wi.import_watchlist("list.csv", csv_text.encode())
    fortios = _asset(watchlist, "fortinet-fortios")
    assert "fortinet:fortios" in fortios["cpe"]
    assert fortios["criticality"] == "critical"
    assert {d["owner"] for d in fortios["deployments"]} == {"Internal", "Customer A"}
    assert next(d for d in fortios["deployments"] if d["owner"] == "Internal")["version"] == "7.4.8"
    assert "import_aliases" not in fortios
    assert report["matched_catalog"] == 1
    assert report["unmatched_products"] == []


def test_unknown_product_becomes_custom_asset_with_overrides():
    csv_text = (
        "owner,product,vendor,version,criticality,category,note,aliases,cpe,kev\n"
        "Customer B,Aplikasi Internal X,Vendor X,2.1.0,tinggi,Web App,,AppX;Aplikasi X,vendorx:appx,Vendor X|AppX\n"
    )
    watchlist, report = wi.import_watchlist("list.csv", csv_text.encode())
    asset = watchlist["assets"][0]
    assert asset["id"].startswith("custom-")
    assert asset["criticality"] == "high"
    assert {"Aplikasi Internal X", "AppX", "Aplikasi X"} <= set(asset["aliases"])
    assert asset["cpe"] == ["vendorx:appx"]
    assert asset["kev"] == [{"vendor": "Vendor X", "product": "AppX"}]
    assert report["unmatched_products"] == ["Aplikasi Internal X"]


def test_xlsx_sheet_name_is_owner_and_model_column_is_not_a_version():
    content = _xlsx({
        "Customer A": [
            ["No", "Category", "Function", "Brand", "Model/Version"],
            [1, "Appliance", "Firewall", "Fortinet", "Fortigate 201E"],
            [2, "Software", "VPN Client", "Forticlient VPN", "7.4.3"],
        ],
        "Internal": [
            ["No", "Kategori", "Tools"],
            [1, "Aplikasi", "Google Chrome"],
        ],
    })
    watchlist, report = wi.import_watchlist("customers.xlsx", content)
    fortios = _asset(watchlist, "fortinet-fortios")
    assert fortios["deployments"] == [{"owner": "Customer A", "note": "Fortinet; Fortigate 201E"}]
    forticlient = _asset(watchlist, "fortinet-forticlient")
    assert forticlient["deployments"][0]["version"] == "7.4.3"
    assert _asset(watchlist, "google-chrome")["deployments"][0]["owner"] == "Internal"
    assert report["owners"] == ["Customer A", "Internal"]


def test_short_aliases_need_word_boundaries():
    catalog = wi.load_catalog()
    assert wi.match_catalog("Microsoft Visio Diagram design", catalog)["id"] == "microsoft-office"
    assert wi.match_catalog("Diagram tool", catalog) is None
    assert wi.match_catalog("Sangfor IAG 13.0", catalog)["id"] == "sangfor"


def test_json_watchlist_is_validated():
    with pytest.raises(wi.WatchlistImportError):
        wi.import_watchlist("bad.json", b'{"items": []}')
    data = {"assets": [{"id": "x", "product": "X", "deployments": [{"owner": "A", "version": "1.0"}]}]}
    watchlist, report = wi.import_watchlist("ok.json", json.dumps(data).encode())
    assert watchlist == data
    assert report["owners"] == ["A"] and report["with_version"] == 1


def test_rejects_unknown_format_and_empty_sheet():
    with pytest.raises(wi.WatchlistImportError):
        wi.import_watchlist("list.txt", b"hello")
    with pytest.raises(wi.WatchlistImportError):
        wi.import_watchlist("empty.csv", b"foo,bar\n1,2\n")


def test_save_keeps_backup(tmp_path, monkeypatch):
    target = tmp_path / "watchlist.json"
    monkeypatch.setattr(wi, "WATCHLIST_FILE", target)
    assert wi.save_watchlist({"assets": []}) is None
    backup = wi.save_watchlist({"assets": [{"id": "a", "product": "A", "deployments": []}]})
    assert backup and (tmp_path / backup).exists()
    assert json.loads(target.read_text(encoding="utf-8"))["assets"][0]["id"] == "a"


def test_export_round_trips_through_import():
    csv_text = "owner,product,version\nInternal,FortiGate,7.4.8\nCustomer A,Google Chrome,\n"
    watchlist, _ = wi.import_watchlist("list.csv", csv_text.encode())
    exported = wi.export_watchlist_xlsx(watchlist["assets"])
    assert load_workbook(io.BytesIO(exported)).sheetnames[0] == "Watchlist"
    again, report = wi.import_watchlist("export.xlsx", exported)
    assert {a["id"] for a in again["assets"]} == {a["id"] for a in watchlist["assets"]}
    assert report["with_version"] == 1


def test_template_xlsx_is_importable():
    watchlist, report = wi.import_watchlist("template.xlsx", wi.template_xlsx())
    assert report["owners"] == ["Customer A", "Customer B", "Internal"]
