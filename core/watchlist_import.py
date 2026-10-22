"""
Watchlist import/export for LazySOC.

Turns an uploaded critical tools list (LazySOC template .xlsx/.csv, a per-customer .xlsx where each sheet
is one owner, or a watchlist .json) into data/watchlist.json. Rows are enriched from the generic product
catalog (core/product_catalog.json) so aliases, CPE and KEV names - needed for version checks - do not
have to be typed by hand.
"""

import csv
import io
import json
import re
import shutil
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from core.config import WATCHLIST_FILE

CATALOG_FILE = Path(__file__).resolve().parent / "product_catalog.json"
MAX_ROWS = 5000
TEMPLATE_COLUMNS = ["owner", "product", "vendor", "version", "criticality", "category", "note", "aliases", "cpe", "kev"]

# Header synonyms (lowercased, whitespace-collapsed). Per-customer sheets often use Indonesian or vendor wording.
HEADER_SYNONYMS = {
    "owner": {"owner", "customer", "pelanggan", "pemilik", "client"},
    "product": {"product", "produk", "tools", "tool", "tools/apps", "produk/brand", "brand", "aplikasi", "application", "software"},
    "vendor": {"vendor", "manufacturer"},
    "version": {"version", "versi", "version os", "model/version", "model / version", "firmware"},
    "criticality": {"criticality", "kritis", "tingkat kritis", "priority", "prioritas"},
    "category": {"category", "kategori", "jenis", "function", "fungsi"},
    "note": {"note", "notes", "catatan", "hostname", "model", "used by team", "user", "department", "deparement"},
    "aliases": {"aliases", "alias"},
    "cpe": {"cpe"},
    "kev": {"kev"},
}
CRITICALITY_MAP = {
    "critical": "critical", "kritis": "critical", "crit": "critical",
    "high": "high", "tinggi": "high",
    "normal": "normal", "medium": "normal", "sedang": "normal", "low": "normal", "rendah": "normal",
}
RE_VERSION_VALUE = re.compile(r"^v?\d+(?:\.\d+)+\b")


class WatchlistImportError(ValueError):
    pass


def _norm(text: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(text or "").lower())


def _clean(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def _split_list(value: str) -> List[str]:
    return [v.strip() for v in re.split(r"[;\n]", value or "") if v.strip()]


def load_catalog() -> List[Dict[str, Any]]:
    try:
        with open(CATALOG_FILE, "r", encoding="utf-8") as f:
            return json.load(f).get("products", [])
    except (OSError, json.JSONDecodeError):
        return []


def match_catalog(text: str, catalog: List[Dict[str, Any]]) -> Optional[Dict[str, Any]]:
    """Finds the catalog product whose longest alias/product name appears in the row text."""
    haystack = _norm(text)
    lowered = (text or "").lower()
    best, best_len = None, 0
    for entry in catalog:
        names = ([entry.get("product", "")] + entry.get("aliases", []) + entry.get("strict_aliases", [])
                 + entry.get("import_aliases", []))
        for name in names:
            key = _norm(name)
            if len(key) < 3 or len(key) <= best_len:
                continue
            if len(key) < 5:
                # Short names (IAG, EBS, PHP) must stand alone, or "diagram" would match IAG.
                found = re.search(rf"(?<![a-z0-9]){re.escape(name.lower())}(?![a-z0-9])", lowered) is not None
            else:
                found = key in haystack  # tolerant of spacing/punctuation: "Forti Client" == "FortiClient"
            if found:
                best, best_len = entry, len(key)
    return best


# ------------------------------------------------------------------------------
# Row extraction
# ------------------------------------------------------------------------------

def _map_headers(header_cells: List[Any]) -> Dict[str, int]:
    mapping: Dict[str, int] = {}
    for idx, cell in enumerate(header_cells):
        key = _clean(cell).lower()
        for field, synonyms in HEADER_SYNONYMS.items():
            if key in synonyms and field not in mapping:
                mapping[field] = idx
                break
    return mapping


def _rows_from_table(table: List[List[Any]], default_owner: str = "") -> List[Dict[str, str]]:
    """Finds the header row (first row with a product column) and converts the rows below it."""
    header_idx, mapping = None, {}
    for i, row in enumerate(table[:15]):
        candidate = _map_headers(row)
        if "product" in candidate:
            header_idx, mapping = i, candidate
            break
    if header_idx is None:
        return []
    if "owner" not in mapping and not default_owner:
        return []

    rows = []
    for raw in table[header_idx + 1:]:
        def cell(field: str) -> str:
            idx = mapping.get(field)
            return _clean(raw[idx]) if idx is not None and idx < len(raw) else ""

        product = cell("product")
        if not product or product.lower() in ("no", "-"):
            continue
        row = {field: cell(field) for field in TEMPLATE_COLUMNS if field != "owner"}
        row["owner"] = cell("owner") or default_owner
        row["product"] = product

        # "Model/Version" style columns hold either a version (7.4.3) or a model name (FortiGate 201E).
        if row["version"] and not RE_VERSION_VALUE.match(row["version"]):
            row["note"] = "; ".join(v for v in (row["version"], row["note"]) if v)
            row["version"] = ""
        if row["owner"]:
            rows.append(row)
    return rows


def rows_from_csv(content: bytes) -> List[Dict[str, str]]:
    text = content.decode("utf-8-sig", errors="replace")
    table = list(csv.reader(io.StringIO(text)))
    return _rows_from_table(table)


def rows_from_xlsx(content: bytes) -> List[Dict[str, str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise WatchlistImportError("Paket 'openpyxl' belum terpasang (pip install -r requirements.txt).") from e

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as e:
        raise WatchlistImportError(f"File Excel tidak dapat dibaca: {e}") from e

    rows: List[Dict[str, str]] = []
    for ws in wb.worksheets:
        table = [list(r) for r in ws.iter_rows(values_only=True)]
        # Without an owner column, the sheet name is the owner (one sheet per customer).
        rows.extend(_rows_from_table(table, default_owner=_clean(ws.title)))
    wb.close()
    return rows


# ------------------------------------------------------------------------------
# Watchlist building
# ------------------------------------------------------------------------------

def _slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")[:60] or "asset"


def build_watchlist(rows: List[Dict[str, str]], source: str) -> Tuple[Dict[str, Any], Dict[str, Any]]:
    """Groups rows into assets (one per product), enriching known products from the catalog."""
    if not rows:
        raise WatchlistImportError("Tidak ada baris produk yang terbaca. Pastikan ada kolom 'product' (dan 'owner', atau satu sheet per owner).")
    if len(rows) > MAX_ROWS:
        raise WatchlistImportError(f"Terlalu banyak baris ({len(rows)}); maksimal {MAX_ROWS}.")

    catalog = load_catalog()
    assets: Dict[str, Dict[str, Any]] = {}
    unmatched: List[str] = []

    for row in rows:
        match_text = " ".join([row.get("vendor", ""), row["product"], row.get("note", ""), row.get("category", "")])
        entry = match_catalog(match_text, catalog)
        if entry:
            asset_id = entry["id"]
            base = {k: (list(v) if isinstance(v, list) else v) for k, v in entry.items() if k != "import_aliases"}
        else:
            asset_id = "custom-" + _slug(f"{row.get('vendor', '')} {row['product']}")
            base = {
                "id": asset_id,
                "vendor": row.get("vendor") or row["product"],
                "product": row["product"],
                "category": row.get("category", ""),
                "criticality": "normal",
                "aliases": [row["product"]] if len(_norm(row["product"])) >= 3 else [],
            }
            if row["product"] not in unmatched:
                unmatched.append(row["product"])

        asset = assets.setdefault(asset_id, {**base, "deployments": []})

        crit = CRITICALITY_MAP.get(row.get("criticality", "").lower())
        if crit:
            asset["criticality"] = crit
        for alias in _split_list(row.get("aliases", "")):
            if alias not in asset.setdefault("aliases", []):
                asset["aliases"].append(alias)
        for cpe in _split_list(row.get("cpe", "")):
            if cpe.lower() not in [c.lower() for c in asset.setdefault("cpe", [])]:
                asset["cpe"].append(cpe.lower())
        for spec in _split_list(row.get("kev", "")):
            vendor, _, product = spec.partition("|")
            asset.setdefault("kev", []).append({"vendor": vendor.strip(), "product": product.strip()})

        deployment = {"owner": row["owner"]}
        if row.get("version"):
            deployment["version"] = row["version"]
        note = "; ".join(v for v in (row["product"] if entry else "", row.get("note", "")) if v)
        if note:
            deployment["note"] = note
        if deployment not in asset["deployments"]:
            asset["deployments"].append(deployment)

    watchlist = {
        "version": 1,
        "source": f"{source} (imported {datetime.now().strftime('%Y-%m-%d %H:%M')})",
        "assets": list(assets.values()),
    }
    report = {
        "rows": len(rows),
        "assets": len(assets),
        "matched_catalog": sum(1 for a in assets.values() if not a["id"].startswith("custom-")),
        "unmatched_products": unmatched,
        "owners": sorted({d["owner"] for a in assets.values() for d in a["deployments"]}),
        "with_version": sum(1 for a in assets.values() for d in a["deployments"] if d.get("version")),
    }
    return watchlist, report


def _validate_watchlist_json(data: Any) -> Dict[str, Any]:
    if not isinstance(data, dict) or not isinstance(data.get("assets"), list):
        raise WatchlistImportError("JSON harus berformat watchlist LazySOC: {\"assets\": [...]} (lihat watchlist.example.json).")
    for i, asset in enumerate(data["assets"]):
        if not isinstance(asset, dict) or not asset.get("id") or not asset.get("product"):
            raise WatchlistImportError(f"Asset #{i + 1} wajib punya 'id' dan 'product'.")
        if not isinstance(asset.get("deployments", []), list):
            raise WatchlistImportError(f"Asset '{asset['id']}': 'deployments' harus berupa list.")
    return data


def import_watchlist(filename: str, content: bytes) -> Tuple[Dict[str, Any], Dict[str, Any]]:
    """Parses an uploaded file into a watchlist dict plus an import report (nothing is saved yet)."""
    name = (filename or "").lower()
    if name.endswith(".json"):
        try:
            data = _validate_watchlist_json(json.loads(content.decode("utf-8-sig")))
        except json.JSONDecodeError as e:
            raise WatchlistImportError(f"JSON tidak valid: {e}") from e
        report = {
            "rows": sum(len(a.get("deployments", [])) for a in data["assets"]),
            "assets": len(data["assets"]),
            "matched_catalog": None,
            "unmatched_products": [],
            "owners": sorted({d.get("owner", "") for a in data["assets"] for d in a.get("deployments", []) if d.get("owner")}),
            "with_version": sum(1 for a in data["assets"] for d in a.get("deployments", []) if d.get("version")),
        }
        return data, report
    if name.endswith(".csv"):
        return build_watchlist(rows_from_csv(content), filename)
    if name.endswith(".xlsx"):
        return build_watchlist(rows_from_xlsx(content), filename)
    raise WatchlistImportError("Format tidak didukung. Gunakan .xlsx, .csv, atau .json.")


def read_current_watchlist() -> Dict[str, Any]:
    try:
        with open(WATCHLIST_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    except (OSError, json.JSONDecodeError):
        return {"version": 1, "assets": []}


def save_watchlist(watchlist: Dict[str, Any]) -> Optional[str]:
    """Writes data/watchlist.json, keeping a timestamped backup of the previous file."""
    backup_name = None
    WATCHLIST_FILE.parent.mkdir(parents=True, exist_ok=True)
    if WATCHLIST_FILE.exists():
        backup = WATCHLIST_FILE.with_name(f"watchlist.backup-{datetime.now().strftime('%Y%m%d-%H%M%S')}.json")
        shutil.copy2(WATCHLIST_FILE, backup)
        backup_name = backup.name
    tmp = WATCHLIST_FILE.with_suffix(".json.tmp")
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(watchlist, f, indent=2, ensure_ascii=False)
        f.write("\n")
    tmp.replace(WATCHLIST_FILE)
    return backup_name


# ------------------------------------------------------------------------------
# Export (template / current list) as .xlsx
# ------------------------------------------------------------------------------

TEMPLATE_EXAMPLES = [
    ["Internal", "FortiGate", "Fortinet", "7.4.8", "critical", "Firewall", "edge-fw-01", "", "", ""],
    ["Customer A", "FortiClient VPN", "Fortinet", "7.4.3", "", "VPN Client", "", "", "", ""],
    ["Customer A", "Google Chrome", "Google", "", "", "Browser", "", "", "", ""],
    ["Customer B", "Aplikasi Internal X", "Vendor X", "2.1.0", "high", "Web App", "", "AppX;Aplikasi X", "", ""],
]


def _xlsx_bytes(rows: List[List[Any]], title: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append(TEMPLATE_COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    for col, width in zip("ABCDEFGHIJ", (16, 30, 18, 14, 12, 18, 30, 24, 28, 24)):
        ws.column_dimensions[col].width = width

    help_ws = wb.create_sheet("Petunjuk")
    for line in (
        ["Kolom wajib: owner, product. Kolom lain opsional."],
        ["version: isi versi nyata (mis. 7.4.8) agar cek versi CVE/KEV bisa jalan."],
        ["criticality: critical / high / normal. Kosong = default dari katalog produk."],
        ["aliases, cpe: pisahkan dengan titik koma. kev: vendor|product, mis. Fortinet|FortiOS."],
        ["Produk terkenal (FortiGate, Chrome, VMware, dll.) otomatis dilengkapi alias/CPE/KEV dari katalog."],
    ):
        help_ws.append(line)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def template_xlsx() -> bytes:
    return _xlsx_bytes(TEMPLATE_EXAMPLES, "Watchlist")


def export_watchlist_xlsx(watchlist_assets: List[Dict[str, Any]]) -> bytes:
    rows = []
    for asset in watchlist_assets:
        kev = ";".join(f"{k.get('vendor', '')}|{k.get('product', '')}" for k in asset.get("kev", []) or [])
        for d in asset.get("deployments", []):
            rows.append([
                d.get("owner", ""), asset.get("product", ""), asset.get("vendor", ""), d.get("version", ""),
                asset.get("criticality", ""), asset.get("category", ""), d.get("note", ""),
                ";".join(asset.get("aliases", []) or []), ";".join(asset.get("cpe", []) or []), kev,
            ])
    return _xlsx_bytes(rows, "Watchlist")
